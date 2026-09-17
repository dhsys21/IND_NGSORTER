from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from hashlib import sha256
import csv, io, json
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
# Reuse only the verified source-template component functions, not its content.
prefix = (HERE/'build_maintenance.py').read_text(encoding='utf-8').split('cover = [')[0]
exec(compile(prefix, str(HERE/'build_maintenance.py'), 'exec'))
TMP = HERE/'rev002'
TMP.mkdir(exist_ok=True)
REF = BASE/'01.Document/설비매뉴얼/NGSORTER_96CH_India_Maintenance_Manual_KR_20260917 001.docx'
OUT = TMP/'NGSORTER_96CH_India_Maintenance_Manual_KR_20260917 002.docx'
raw = REF.read_bytes()
assert sha256(raw).hexdigest() == 'ea19561955fc0321f898482bd584835b6491034167ef606e9a43b404f0ce8ba5'
original = ZipFile(io.BytesIO(raw))
doc = Document(REF)
body = doc._element.body
cursor = None
def append(el):
    global cursor
    if cursor is None: body.insert(len(body)-1,el)
    else:
        cursor.addnext(el)
        cursor=el

def textof(el): return ''.join(el.itertext()) if False else ''.join(el.xpath('.//w:t/text()'))
def heading(text):
    return next(p for p in doc.paragraphs if p.text==text and p.style.style_id=='10')
def replace_section(text):
    global cursor
    el=heading(text)._p
    cursor=el.getprevious()
    while el is not None and el.tag!=qn('w:sectPr'):
        nxt=el.getnext()
        body.remove(el)
        if nxt is not None and nxt.tag==qn('w:p'):
            pstyle=nxt.find('./'+qn('w:pPr')+'/'+qn('w:pStyle'))
            if pstyle is not None and pstyle.get(qn('w:val'))=='10':break
        el=nxt
def insert_before(text):
    global cursor
    cursor=heading(text)._p.getprevious()
def replace_text(old,new):
    for par in doc.paragraphs:
        if par.text==old:
            for run in par.runs:run.text=''
            par.add_run(new)
            return
    raise ValueError(old)

replace_text('문서번호: NGS-MM-96CH-001    Rev. 0.1','문서번호: NGS-MM-96CH-001    Rev. 0.2')
history=next(t for t in doc.tables if t.cell(0,0).text=='Rev.')
row=deepcopy(history.rows[-1]._tr)
for tc,value in zip(row.findall(qn('w:tc')),['0.2','2026.09.17','8/13 서보 백업값, 3축 설정, HOME 속도, Z 속도 화면 및 Device Monitor 보완','대흥시스텍']):
    texts=tc.xpath('.//w:t');texts[0].text=value
    for tx in texts[1:]:tx.text=''
history._tbl.append(row)

backup=MAN/'IND_NGSORTER_SERVO 20260813'
params={(int(a),int(b,16)):int(c,16) for a,b,c in csv.reader((backup/'SampleData.prm2').read_text().splitlines())}
def hx(axis,number): return f'{params[axis,number]:04X}h'
def paramrow(number,label):return [f'{number:04X} {label}',*[hx(a,number) for a in (1,2,3)]]

replace_section('3.2 프로젝트와 3축 구성')
title('3.2 프로젝트와 3축 구성')
p('8월 13일 백업은 Axis1=X, Axis2=Y, Axis3=Z의 독립 3축입니다. 축 사용 여부는 System의 총 축수 항목이 아니라 Parameter Data → SampleData.prm2 → Control에서 축별로 확인합니다.')
table(['Control 항목','Axis1 X','Axis2 Y','Axis3 Z'],[
 paramrow(0x200,'OPC1'),paramrow(0x203,'AXALC'),paramrow(0x264,'TGRP'),paramrow(0x265,'TOP')],[3100,1990,1990,1991])
steps(['백업 폴더의 IND_NGSORTER_SERVO.pbp2를 작업용 사본으로 열고, Parameter Data의 SampleData.prm2를 엽니다. Point Data(.pnt2)와 구분합니다.',
 'Control 0200 OPC1의 사용 축 설정을 확인합니다. 이 백업의 Axis1~3은 0001h, Axis4~32는 0000h(미사용)입니다. 다른 옵션이 섞인 프로젝트에는 이 전체 Word 값을 무조건 덮어쓰지 않습니다.',
 '0203 AXALC 값 0001h/0002h/0003h 및 앰프 실물 축 번호를 대조합니다. 현장 설정 이력은 스위치 0/1/2, 앰프 표시 n01/n02/n03입니다.',
 '0264 TGRP와 0265 TOP이 세 축 모두 0000h인지 확인합니다. 승인 설정을 적용한 뒤 Positioning Test의 Axis 목록에 Axis1, Axis2, Axis3가 각각 표시되는지 확인합니다.'])
fig(PB/'servo setting 13.png','그림 3-2. 현장 화면 이력: 독립된 Axis1, Axis2, Axis3 선택',2.65)
p('백업의 System 0002 SYSOP2는 0000h로 축번호 자동 할당 상태입니다. 0203 값을 기록했다고 수동 할당 모드가 활성화된 것은 아닙니다. 이 구성에서는 앰프 번호와 보드 축의 기본 대응도 함께 확인하며, SYSOP2를 임의 변경하지 않습니다.')

replace_section('3.3 System startup과 파라미터 저장')
title('3.3 System startup과 파라미터 저장')
fig(PB/'servo setting 4.png','그림 3-3. Positioning Test에서 Parameter 선택 후 System startup',2.9)
steps(['기존 보드값을 읽어 백업하고, Positioning Test의 Parameter에 승인된 SampleData.prm2를 선택합니다. Edit로 축·센서·전자기어·원점값을 확인합니다.',
 'Do not send servo parameters는 앰프 파라미터 전송 여부입니다. 복구 범위와 제조사 적용 조건에 맞게 사용하며, 이름만 보고 임의 선택하지 않습니다.',
 '기계 정지 및 안전 확보 후 System startup을 실행합니다. 적용값을 읽어 확인하고 축 목록과 알람, Ready를 확인합니다.',
 'PC의 .prm2 저장과 보드 Flash ROM 보존은 별개입니다. 승인된 백업 절차에 따라 각각 확인하고, 변경 후 HOME·방향·정량 이동을 재검증합니다.'])
sub('운영 프로그램 OPEN과 파일 적용')
p('확인한 현재 코드의 신규 서보 기동은 D:\\IND_NGSORTER_SERVO\\SampleData.prm2를 읽습니다. System과 Axis1~3 값을 검증·적용하고 System start를 실행하며, 이 경로는 Flash ROM을 변경하지 않습니다.')
p('이미 보드가 RUNNING이면 기존 시스템을 재사용하여 재기동과 파일 재적용을 생략합니다. 따라서 파일만 수정한 뒤 OPEN을 누르거나 프로그램만 다시 실행했다고 새 값이 적용된 것으로 판단하면 안 됩니다. 승인된 정지·재기동 절차 후 실제 보드값을 읽어 확인합니다.')
note('8/13 백업 폴더와 운영용 D:\\IND_NGSORTER_SERVO 폴더는 다른 경로입니다. 변경·배포·복원 파일을 혼동하지 않습니다. 운영 PC의 설치 버전이 이 코드와 같은지도 확인합니다.')

replace_section('3.5 전자기어와 1 mm 단위 확인')
title('3.5 전자기어와 1 mm 단위 확인')
p('다음은 8월 13일 SampleData.prm2에서 직접 읽은 3축 값입니다. 과거 4축 화면의 5974h, 1388h 등과 구분합니다. h는 16진수이며 괄호는 10진수입니다.')
table(['파라미터','Axis1 X','Axis2 Y','Axis3 Z'],[
 paramrow(0x20a,'CMXL 하위'),paramrow(0x20b,'CMXH 상위'),
 paramrow(0x20c,'CDVL 하위'),paramrow(0x20d,'CDVH 상위'),
 paramrow(0x20e,'SUML 하위'),paramrow(0x20f,'SUMH 상위'),
 ['합성 분자 N','4,194,304','4,194,304','4,194,304'],
 ['합성 분모 D','D084h (53,380)','4E20h (20,000)','2710h (10,000)'],
 ['합성 속도 단위 배율','4E20h (20,000)','4E20h (20,000)','2710h (10,000)']],[3100,1990,1990,1991])
p('32비트 값은 하위 Word + 상위 Word × 65,536입니다. 예를 들어 분자 0040h:0000h는 4,194,304입니다. 두 Word를 하나의 값으로 대조합니다.')
note('X축은 분모 53,380과 속도 단위 배율 20,000이 서로 다릅니다. 백업에 실제로 이렇게 저장되어 있으므로 두 값을 임의로 같게 맞추지 않습니다.')
sub('1000 지령을 1 mm로 검증')
steps(['실제 모터·엔코더 및 감속기·볼스크류·벨트의 1회전 이동량과 승인 티칭 단위를 확인합니다. MR-J4-100B/70B/20B는 앰프 형명입니다.',
 '보드에 적용된 020A~020F를 읽어 이 표 또는 승인된 최신값과 대조합니다. 파일값만으로 실측 단위가 맞는다고 확정하지 않습니다.',
 '3.6절의 INC 저속 시험으로 10진 1000 지령=1 mm, 10000 지령=10 mm를 각 축에서 계측합니다. 정·역방향과 반복성을 확인합니다.',
 '위치 기어를 바꾸면 티칭·소프트리밋·HOME·속도 단위도 재검증합니다. 기존 티칭으로 즉시 생산하지 않습니다.'])
p('동일 엔코더라도 기계의 1회전 이동량이 다르면 기어비는 다릅니다. SUML/SUMH는 속도 단위 항목이며, 위치 오차를 고치기 위해 단독 변경하는 항목이 아닙니다.')

insert_before('4.1 CC-Link 설치와 Master 설정')
title('3.7 8월 13일 백업과 설정값 변경')
p('확인 원본: 03.Manual\\IND_NGSORTER_SERVO 20260813 폴더. 원본을 보존하고 변경일·설비명을 붙인 사본에서 작업합니다. 파일명을 바꾼 경우 PB Test의 프로젝트 참조와 운영용 배포 파일도 확인합니다.')
table(['파일 또는 항목','백업에서 확인한 내용'],[
 ['IND_NGSORTER_SERVO.pbp2','MR-MC240 / Board ID 0 / Channel 1. Parameter 참조: SampleData.prm2'],
 ['SampleData.prm2','System, Servo, Control 등 실제 파라미터. 축별 값은 3.2·3.5·3.8절 참조'],
 ['SampleData.pnt2','PB Test 위치 시험용 Point Data. 운영 프로그램의 96채널 티칭 파일과 별개'],
 ['*.adr2','Memory Map 화면 구성. 파라미터 백업(.prm2)과 구분'],
 ['Servo 110D POL','Axis1=0000h / Axis2=0000h / Axis3=0001h. 방향은 실물 검증 필요'],
 ['운영용 파일','현재 소스의 읽기 위치: D:\\IND_NGSORTER_SERVO\\SampleData.prm2']])
sub('실제 값을 변경하는 순서')
steps(['NGSORTER와 PB Test가 동시에 보드를 제어하지 않도록 하고, 기존 정상 파일과 현재 보드값을 별도로 백업합니다.',
 'PB Test에서 작업용 프로젝트를 열고 Parameter Data → SampleData.prm2의 해당 탭·번호·Axis 열을 선택합니다. 16진 표시인지 확인하고 변경 전 값을 기록합니다.',
 '축 사용은 Control 0200, 축 대응은 0203, 독립 구동은 0264/0265, HOME 속도는 0242/0243과 024C, 전자기어는 020A~020F를 각 설명에 따라 변경합니다.',
 '파일의 변경값을 저장합니다. 변경 항목에 필요한 보드 적용과 System startup을 수행하고, 보드에서 다시 읽은 값으로 적용 여부를 확인합니다. 읽기 작업으로 미저장 편집값이 덮이지 않게 주의합니다.',
 '단위·방향·HOME 및 저속 동작을 검증한 파일만 운영용 위치로 배포합니다. 운영용 기존 파일은 먼저 보관하고, 배포 경로와 날짜를 작업 기록에 남깁니다.',
 '운영 프로그램 신규 기동 또는 기존 RUNNING 재사용 여부를 로그에서 확인합니다. 파일 저장, 보드 적용, Flash ROM 보존은 각각 확인합니다.'])
note('제조사 설정용 항목과 안전 입력 관련 항목은 일반 속도·축 설정 작업에서 변경하지 않습니다. 8/13 백업은 이력 기준이며 이후 승인된 현장 수정이 있으면 최신 백업이 우선합니다.')

title('3.8 HOME 속도 설정과 적용 위치')
p('HOME 속도는 PB Test의 Parameter Data → SampleData.prm2 → Control에서 설정합니다. 운영 프로그램의 Jog/Z 속도 창은 HOME 속도를 바꾸지 않습니다.')
table(['Control 항목','8/13 백업 X/Y/Z 공통','의미'],[
 ['0240 OPZ1','0018h','원점복귀 방식·방향 등 옵션. 승인된 기존 설정 유지'],
 ['0242 ZSPL / 0243 ZSPH','0032h / 0000h → 50','원점복귀 속도 하위 / 상위'],
 ['0244 ZTCA','1000h → 4,096 ms','원점복귀 가속 시정수'],
 ['0245 ZTCD','1000h → 4,096 ms','원점복귀 감속 시정수'],
 ['024C CRF','0014h → 20','DOG 검출 후 접근(Creep) 속도'],
 ['0246 ZPSL / 0247 ZPSH','0000h / 0000h','원점복귀 완료 좌표']],[2850,2700,3521])
note('50과 20은 보드 Speed units입니다. 50 mm/s, 20 mm/s로 단정하지 않습니다. 0200의 속도 단위와 020E/020F 배율 및 실제 기계 단위를 함께 확인합니다. 1000h는 1,000이 아니라 4,096입니다.')
sub('변경과 시험')
steps(['축별 HOME 경로와 DOG·리밋 극성이 검증된 상태에서 작업용 .prm2를 엽니다.',
 '변경하려는 축의 0242/0243을 조정하고 필요 시 024C 및 0244/0245도 승인값으로 조정합니다. 다른 축 열과 제조사 설정은 건드리지 않습니다.',
 '3.3절에 따라 저장·적용·읽기 확인합니다. PB Test의 ZRN에서 개별 축 시험을 하고, 운영 프로그램에서는 HOME 순서를 확인합니다.',
 '현재 코드의 HOME은 sscHomeReturnStart()로 기동하며 이 호출에서 별도 속도를 전달하지 않습니다. 보드에 적용된 원점 파라미터를 사용하고 Z 완료 후 X/Y를 진행합니다.'])
p('PB Test에서만 값을 바꾸고 운영용 .prm2를 갱신하지 않으면 운영 프로그램의 다음 신규 기동에서 파일값이 다시 적용될 수 있습니다. 반대로 RUNNING 시스템 재사용 중에는 새 파일값이 적용되지 않습니다. 변경한 파일과 보드값을 함께 관리합니다.')
p('근거: Mitsubishi MR-MC200/300 User’s Manual (Details), IB-0300223, 11.3 Control Parameters, 인쇄 쪽 684~685 및 제공된 MR-MC2XX 기동편. 현재 코드: ModRobostar.cpp의 StartServoSystemFromParameterFile / Home.')

insert_before('5.1 연기감지기용 직렬 포트 설치') if any(p.text=='5.1 연기감지기용 직렬 포트 설치' for p in doc.paragraphs) else None
# Locate the first chapter 5 heading, without relying on a translated title.
chapter5=next(p.text for p in doc.paragraphs if p.style.style_id=='10' and p.text.startswith('5.'))
insert_before(chapter5)
title('4.5 CC-Link Device Monitor')
p('CC-Link Ver.2 Utility의 Device Monitor 버튼으로 입력 X와 출력 Y를 확인합니다. 아래는 2026.08.06 운영 PC 캡처이며, 현재 정상 출력 조합을 지시하는 표가 아닙니다.')
fig(PB/'servo on 조건 2026 08 06.png','그림 4-4. Device Monitor: 좌측 X 입력, 우측 Y 출력, Bit device 표시',2.25)
steps(['Board information에서 링크 상태와 오류를 먼저 확인한 후 Device Monitor를 엽니다.',
 'Setting에서 실제 보드 Channel과 표시할 디바이스 범위를 확인합니다. 캡처 하단은 81:CC-Link(1 slot), Network No.0, Station No.255입니다. Station 255는 이 모니터의 접근 설정 표시이므로 실물 Remote 국번으로 옮겨 설정하지 않습니다.',
 'Data Format은 Bit device로 확인합니다. X0000~X003F 및 Y0000~Y003F를 표시하고 현재 할당 주소와 대조합니다.',
 '안전한 점검 조건에서 센서 한 개의 실제 상태 변화와 X 입력 변화, 운영 프로그램 I/O Monitoring의 같은 주소를 비교합니다.',
 'Y는 PC의 출력 명령값 확인에 사용합니다. Y=1만으로 릴레이·실린더가 실제 동작했다고 판단하지 말고 관련 X 피드백과 실물을 확인합니다.'])
table(['관찰 주소','확인할 의미'],[
 ['X0020 / X0021','CHUCK / UNCHUCK 확인'],['X0022 / X0023','셀 감지(OFF=감지, 통신 정상 전제) / BUFFER'],
 ['X0024 / X002D','EMS SWITCH 01 / 02. 과거 캡처 X002D는 현재 의미와 다를 수 있음'],
 ['X002B / X002C','비상정지 측 / 도어 측 안전회로 성립 입력'],
 ['Y0030~Y003E','그리퍼, 안전 Reset, 키락, 램프 등 출력 명령. 4.4절과 현재 I/O 도면 대조']])
note('Device Write로 출력 강제 ON/OFF를 시험하는 절차는 포함하지 않습니다. 생산 프로그램과 Utility가 동시에 출력을 쓰면 충돌할 수 있습니다. 링크 오류·정지된 모니터의 0/1을 최신 센서 상태로 해석하지 않습니다.')

replace_section('9.4 JOG와 Z축 속도 설정')
title('9.4 JOG와 Z축 속도 설정')
dialog=TMP/'captures/jog-z-dialog.png'
if dialog.exists():
    fig(dialog,'그림 9-2. Jog/Z 속도 설정창 실제 캡처. 노트북 표시값은 조작 설명용이며 현장 승인값이 아님',2.65,4.75)
else:
    fig(CAP/'jog.png','그림 9-2. Teaching의 Jog/Z 속도 버튼 위치. 세부 설정창은 현장 프로그램에서 확인',2.1)
table(['항목','허용값 / 기본값','적용'],[
 ['JOG 속도','1~200 / 저장값','버튼을 누르는 방향 JOG'],
 ['Z 고속','300~2700 / 600','Source 전체 하강, Target 감속 시작 Z까지'],
 ['Z 저속','100~500 / 300, 고속 이하','Target 감속 시작 Z 이후 최종 삽입 높이까지'],
 ['Target 감속 시작 Z','0이 아닌 절대 지령 위치','0과 최종 Target Z 사이. 실행 시 현재 Z와 최종 Z 사이인지도 확인']],[2250,2870,3951])
steps(['MANUAL에서 실제 축 정지를 확인하고 Teaching → Jog/Z 속도를 클릭합니다. 축 이동 중에는 설정창이 열리지 않습니다.',
 'JOG, Z 고속, Z 저속 및 Target 감속 시작 Z를 입력합니다. 감속 시작값은 거리에 대한 %가 아니라 티칭과 같은 단위의 절대 좌표입니다.',
 'Save를 누르면 검증 후 설정 파일에 저장되어 다음 기동에도 읽습니다. Cancel은 변경을 적용하지 않습니다. HOME 속도는 이 창과 무관하며 3.8절에서 설정합니다.',
 '빈 트레이 저속 시험에서 Target의 고속 구간 → 실제 정지 확인 → 저속 최종 구간을 검증합니다. Source는 고속 한 구간으로 하강하므로 별도로 확인합니다.'])
p('화면의 Z first 80% speed는 Z 고속, Z final 20% speed는 Z 저속 항목입니다. 옛 명칭과 달리 80%/20% 자동 분할이 아니라 Target Z slow-start absolute position의 절대 높이에서 전환합니다. 기본 600/300보다 저장값이 우선합니다.')
note('감속 시작점이 이동 구간 밖이면 하강이 차단됩니다. 중간 높이에서 정지한 채 이어지지 않으면 [Z SPEED] 로그, 보드 이동 상태와 실제 위치를 확인하고 반복 기동하지 않습니다.')

# Regenerate the static TOC using the final Word pagination, while retaining its template formatting.
toc=heading('목 차')._p
el=toc.getnext()
while el is not None and el!=heading('1. 안전 및 작업 준비')._p:
    nxt=el.getnext();body.remove(el);el=nxt
heads=[p.text for p in doc.paragraphs if p.style.style_id=='10' and p.text not in ('문서 정보','목 차')]
pagefile=TMP/'toc_pages.json'
pagemap=json.loads(pagefile.read_text(encoding='utf-8')) if pagefile.exists() else {}
anchor=toc
for text in heads:
    el=deepcopy(ps[17])
    for c in list(el):
        if c.tag!=qn('w:pPr'):el.remove(c)
    par=Paragraph(el,doc._body)
    par.paragraph_format.space_after=Pt(1)
    par.paragraph_format.space_before=Pt(0)
    par.add_run(text+'\t'+str(pagemap.get(text,'—')))
    anchor.addnext(el);anchor=el
doc.core_properties.revision=2
buf=io.BytesIO();doc.save(buf)
generated=ZipFile(io.BytesIO(buf.getvalue()))
editable={'word/document.xml','word/_rels/document.xml.rels','docProps/core.xml','[Content_Types].xml'}
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
    for name in original.namelist():z.writestr(name,generated.read(name) if name in editable else original.read(name))
    for name in generated.namelist():
        if name not in original.namelist():z.writestr(name,generated.read(name))
with ZipFile(OUT) as z:
    for name in original.namelist():
        if name not in editable:assert z.read(name)==original.read(name),name
assert REF.read_bytes()==raw
(TMP/'headings.json').write_text(json.dumps(heads,ensure_ascii=False,indent=2),encoding='utf-8')
(TMP/'build_evidence.json').write_text(json.dumps({'reference':str(REF),'reference_sha256':sha256(raw).hexdigest(),'backup_sha256':sha256((backup/'SampleData.prm2').read_bytes()).hexdigest(),'headings':len(heads),'actual_dialog_capture':dialog.exists(),'unchanged_parts':'PASS'},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT)
print('PASS original and preserve-only parts unchanged;',len(heads),'headings; dialog:',dialog.exists())
