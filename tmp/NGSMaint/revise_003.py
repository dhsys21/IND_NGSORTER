from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from hashlib import sha256
import io, json, re
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
prefix = (HERE/'build_maintenance.py').read_text(encoding='utf-8').split('cover = [')[0]
exec(compile(prefix, str(HERE/'build_maintenance.py'), 'exec'))
TMP = HERE/'rev003'
TMP.mkdir(exist_ok=True)
REF = BASE/'01.Document/설비매뉴얼/NGSORTER_96CH_India_Maintenance_Manual_KR_20260917 002.docx'
OUT = TMP/'NGSORTER_96CH_India_Maintenance_Manual_KR_20260917 003.docx'
raw = REF.read_bytes()
assert sha256(raw).hexdigest() == 'd76c45d34d5547491216d503913e9424efcf4dc916cec901590d1046a15ea560'
original = ZipFile(io.BytesIO(raw))
doc = Document(REF)
body = doc._element.body
els = list(body)
changes = []

def setpara(par, value):
    # Keep the template paragraph and first run formatting intact.
    if par.runs:
        par.runs[0].text = value
        for run in par.runs[1:]:
            run.text = ''
    else:
        par.add_run(value)

def replace(index, value):
    el = els[index]
    assert el.tag == qn('w:p'), index
    par = Paragraph(el, doc._body)
    changes.append({'index': index, 'before': par.text, 'after': value})
    if value is None:
        body.remove(el)
    else:
        setpara(par, value)

def cellreplace(index, old, new):
    tableobj = Table(els[index], doc._body)
    found = 0
    for row in tableobj.rows:
        for cell in row.cells:
            if cell.text == old:
                setpara(cell.paragraphs[0], new)
                for par in cell.paragraphs[1:]:
                    par._element.getparent().remove(par._element)
                found += 1
    assert found == 1, (index, old, found)
    changes.append({'table_index': index, 'before': old, 'after': new})

replacements = {
4: '문서번호: NGS-MM-96CH-001    Rev. 0.3',
12: '본 문서는 NGSORTER의 설치, 설정 복원, 티칭, 수동 조작 및 예방 정비 절차를 설명합니다. 작업 전 설비 ID와 적용 프로그램 버전을 확인하고 승인된 설정 파일을 백업합니다.',
13: '설정 변경 후에는 적용값을 다시 읽어 확인하고, 센서·원점·정량 이동 및 관련 인터락을 검증합니다. 작업 결과와 변경 전후 설정을 정비 기록으로 보관합니다.',
14: '설치와 변경은 설비를 안전 정지한 상태에서 실시합니다. 안전장치 및 기구 내부 작업은 지정된 담당자가 승인된 절차에 따라 수행합니다.',
93: 'PC 교체 시 장치 식별 정보, COM 번호 및 네트워크 설정을 복원하고 실제 배선과 데이터 수신을 확인합니다.',
98: '4. Project Setting에서 MR-MC240, Board ID 0, Channel 1, Interface BUS를 확인합니다.',
102: '그림 3-1. Project Setting: 모델, Interface, Board ID 및 Channel',
103: 'PC 내부 MR-MC240 보드의 Interface는 BUS를 사용합니다.',
110: '3. 0203 AXALC 값 0001h/0002h/0003h와 앰프 실물 축 번호를 대조합니다. 스위치 0/1/2 및 앰프 표시 n01/n02/n03의 대응을 확인합니다.',
113: '그림 3-2. 독립된 Axis1, Axis2, Axis3 선택',
123: '운영 프로그램의 신규 서보 기동은 D:\\IND_NGSORTER_SERVO\\SampleData.prm2를 읽습니다. System과 Axis1~3 값을 검증·적용하고 System start를 실행하며, 이 경로는 Flash ROM을 변경하지 않습니다.',
125: '백업은 날짜별 폴더에 보관하고, 검증을 완료한 파일을 운영용 D:\\IND_NGSORTER_SERVO 폴더에 배포합니다. 변경·배포·복원 파일을 혼동하지 않습니다.',
138: '다음은 2026.08.13 SampleData.prm2 백업의 3축 전자기어 및 속도 단위 설정값입니다. h는 16진수이며 괄호는 10진수입니다.',
142: 'X축의 위치 분모는 53,380, 속도 단위 배율은 20,000입니다. 위치 기어와 속도 단위 배율은 별도 항목이므로 두 값을 임의로 같게 맞추지 않습니다.',
157: '7. 조정 시 Servo OFF 등 제조사 적용 조건을 지키고 파라미터 재적용·읽기 확인 후 다시 시험합니다. 최종 설정과 측정값을 저장합니다.',
162: '16.3은 Encoder initial communication / Receive data error 3입니다. 전원 차단 후 CN2 엔코더 계통·케이블·모터 연결을 먼저 점검합니다. CN8 등 안전 관련 커넥터를 임의로 단락하지 않습니다. Ax unmounted는 실제 축 번호와 할당·사용 축 설정을 먼저 대조합니다.',
164: '백업 폴더: 03.Manual\\IND_NGSORTER_SERVO 20260813. 원본을 보존하고 변경일·설비명을 붙인 사본에서 작업합니다. 파일명을 바꾼 경우 PB Test의 프로젝트 참조와 운영용 배포 파일도 확인합니다.',
174: '제조사 설정용 항목과 안전 입력 관련 항목은 일반 속도·축 설정 작업에서 변경하지 않습니다.',
184: '4. 운영 프로그램의 HOME은 보드에 적용된 원점 파라미터를 사용하며, Z 원점복귀 완료 후 X/Y를 진행합니다. HOME 실행 시 별도 이동 속도를 지정하지 않습니다.',
186: None,
192: '5. 전송속도 10 Mbps, Remote net Ver.1 mode를 실물 국 설정 및 승인된 .cbd와 대조합니다.',
196: '그림 4-1. Parameter settings의 Standby master station 표시',
197: 'Remote 장치의 국번호는 실물 DIP 스위치와 .cbd 설정을 함께 확인합니다.',
206: '그림 4-3. 데이터 링크 오류와 Timeout 표시',
210: '4. Network Test 결과의 정확한 코드와 Details를 보존합니다. 0xB310 같은 코드만 보고 원인을 단정하지 않고 해당 보드의 오류 코드표와 대조합니다.',
214: '서보 전원을 켜기 위해 Utility에서 Y0032~Y003D를 강제 쓰지 않습니다. 승인된 안전회로 복구 절차를 사용합니다.',
221: '도어 전원·CC-Link가 꺼졌을 때의 OFF를 문 닫힘으로 판정하지 않습니다. 프로그램의 비상정지 정지 판단은 EMS 01을 사용합니다. 같은 극성인 EMS 02를 포함하여 두 입력의 변화와 실제 안전회로를 함께 시험합니다.',
223: 'CC-Link Ver.2 Utility의 Device Monitor 버튼으로 입력 X와 출력 Y를 확인합니다.',
227: '2. Setting에서 실제 보드 Channel과 표시할 디바이스 범위를 확인합니다. Channel 81:CC-Link(1 slot), Network No.0 및 Station No.255 표시를 확인합니다. Station 255는 모니터 접근 설정이며 실물 Remote 국번호가 아닙니다.',
235: '운영 PC에는 Intel AMT SOL이 설치되어 있습니다. 재설치 시 장치 관리자에서 장치 식별 정보와 드라이버 상태를 확인합니다.',
237: '그림 5-1. Intel Active Management Technology SOL 드라이버 선택',
239: '그림 5-2. Intel AMT SOL 설치 결과',
240: '설정 경로는 장치 관리자 → PCI Serial Port → 드라이버 업데이트 → 컴퓨터에서 찾아보기 → 목록에서 직접 선택 → Ports(COM & LPT) → Intel AMT SOL입니다. PC 제조사·하드웨어 ID에 맞는 승인 드라이버 패키지를 사용합니다.',
241: '감지기 배선이 연결된 포트와 정상 데이터 수신을 함께 확인합니다. 다른 시리얼 카드에 Intel AMT SOL 드라이버를 무조건 선택하지 않습니다.',
244: '그림 5-3. 포트 통신속도와 COM 번호 설정',
245: '1. 장치 관리자에서 해당 직렬 장치의 포트 설정을 열어 115200 bps, 8 data bits, None parity, 1 stop bit를 확인합니다.',
246: '2. 고급 설정에서 운영 COM 번호를 설정하고 다른 장치의 포트 번호와 충돌하지 않는지 확인합니다.',
250: '드라이버 호환성 경고가 발생하면 설치를 중지하고 하드웨어 ID와 제조사 지원 여부를 확인합니다.',
265: '그림 6-1. TSD-V50 레지스터 및 Status 비트 표',
272: '그림 6-2. Human Automation 통신 사양과 프레임',
276: '프로그램의 기본값에 의존하지 말고 Baud rate를 115200으로 명시하여 장치 설정과 일치시킵니다.',
279: 'Source와 Target의 바코드 리더를 각각 식별하고 AutoID Navigator로 설정합니다. 교체 시 실물 형명·케이블·설정 파일을 확인합니다.',
282: '3. 검색된 장치의 시리얼 번호를 실물 명판과 대조하여 설정할 리더를 식별합니다.',
286: '그림 7-1. AutoID Navigator 장치 검색 화면',
287: '검색된 장치에 승인된 운영 IP를 설정하고 Source/Target의 장치 명판과 실제 판독 대상을 함께 기록합니다.',
290: '그림 7-2. BCR 네트워크 및 트리거 명령 설정',
296: '그림 7-3. AutoID Navigator 판독 시험',
297: '1. 렌즈·조명·브라켓과 라벨 오염을 점검하고, 트레이의 모따기와 바코드를 지정된 방향으로 정렬합니다.',
308: '앞 주소가 하위 Word, 다음 주소가 상위 Word이며 부호 없는 32비트 정수로 조합합니다. 예: Voltage raw=220000이면 220.000 V입니다. IEEE754 float로 해석하거나 PC에서 다시 1000을 곱하지 않습니다.',
310: '그림 8-1. Interface의 PLC 전력 정보 표시',
314: '8.2 전력 계측 계통 점검',
315: 'PLC 측 전력 계측 계통은 전장 도면과 승인 설정에 따라 점검합니다. 계측기 → PLC → PC → FMS 순서로 값과 단위를 확인하며, PC는 PLC 수신 데이터를 사용합니다.',
317: '1. 계측 위치, 장치 형명·정격, 전압·전류 범위 및 단상/삼상 결선 방식을 확인합니다.',
328: '그림 9-1. Teaching 및 수동 조작 화면',
337: 'Auto calculate의 기준값',
340: '1000=1 mm가 확인된 경우 채널 피치는 45 mm, 후반 기준점 간격은 590 mm입니다. CH12→CH13 사이의 계산 간격은 95 mm입니다.',
356: '그림 9-2. Jog/Z 속도 설정창',
363: 'Z first 80% speed는 Z 고속, Z final 20% speed는 Z 저속 항목입니다. 속도 전환은 이동거리의 80%/20% 자동 분할이 아니라 Target Z slow-start absolute position의 절대 높이를 기준으로 합니다. 기본 600/300보다 저장값이 우선합니다.',
366: 'Teaching 화면 JOG 버튼의 동작 조건은 다음과 같습니다. PB Test의 JOG에는 동일한 프로그램 인터락이 적용되지 않습니다.',
369: '공통 동작 조건',
376: 'Target 채널 선택은 채널 표시 상태와 그리퍼 셀 감지 조건으로 제한됩니다. 점유된 모든 채널을 일괄 차단하는 기능은 아니므로 표시색만 보고 삽입 위치를 판단하지 말고 실제 빈 슬롯 여부와 셀 정렬을 확인합니다.',
379: '전용 수동 Z 하강과 일반 JOG는 채널 이동과 센터링 감시 조건이 다릅니다. 해당 조작에서는 작업자가 트레이 고정과 셀 정렬을 확인하며, 센터링 신호와 실제 상태가 일치할 때만 하강합니다.',
391: 'X0023 BUFFER가 감지되면 프로그램은 조건부 Z 상승 복구 경로를 사용합니다. 이는 모든 충돌에서 안전하게 빠져나온다는 보장이 아닙니다. 셀의 걸림·변형·파손 위험이 있거나 예상 밖 움직임이면 비상정지합니다.',
400: '그림 11-1. 도어 및 비상정지 알람 화면',
411: 'Gateway XML 설정 경로는 D:\\OpcUaGateway_IND\\CONFIG입니다. NGS 장비 선택과 실제 사용 XML을 확인합니다. OPC UA 구조의 NGS.F1NGS01과 장비 JSON 표준키 F1NGS01을 혼동하지 않습니다. 연결 주소·포트는 운영 PC의 승인 설정을 사용합니다.',
417: '청소·윤활·교체 작업은 승인된 보전 계획에 따라 시행합니다. 지정되지 않은 오일이나 세정제를 사용하지 않습니다.',
434: 'BCR의 명판·시리얼 번호·IP·실제 판독 대상을 Source와 Target로 구분하여 기록하고, 정상 설정 파일과 판독 시험 결과를 함께 보관합니다.',
}
for index, value in replacements.items():
    replace(index, value)

cell_edits = [
(9, '2026.09.17 확인 소스 및 제공된 실제 운영 PC 캡처\nNGSORTER 운영 매뉴얼 Rev. 0.2와 같은 양식', 'NGSORTER 설치·설정 및 유지보수 절차'),
(9, '작성 기준', '문서 기준'),
(9, '운영 PC의 승인된 설정을 기준으로 복구합니다. 노트북의 COM 번호·IP·드라이버를 그대로 적용하지 않습니다.', '설비의 승인된 설정을 기준으로 복구하고, 변경 후 안전·동작 검증을 완료합니다.'),
(90, '제공 자료 내 위치', '설치 파일 위치'),
(90, '서보 설정 이력', '서보 설정 백업'),
(90, '03.Manual\\IND_NGSORTER_SERVO 20260813\n과거 백업이므로 현재 승인 설정과 먼저 비교', '03.Manual\\IND_NGSORTER_SERVO 20260813'),
(128, '현장 센서는 감지 시 ON으로 확인된 이력', '감지 시 ON, 비감지 시 OFF'),
(128, '리밋 센서는 비감지 시 ON, 감지 시 OFF 이력', '비감지 시 ON, 감지 시 OFF'),
(165, '백업에서 확인한 내용', '파일 및 설정 내용'),
(165, '현재 소스의 읽기 위치: D:\\IND_NGSORTER_SERVO\\SampleData.prm2', '운영용 읽기 위치: D:\\IND_NGSORTER_SERVO\\SampleData.prm2'),
(201, 'AJ65SBTB1 해당 B RATE의 4=ON, 2=OFF, 1=OFF 설정을 해당 모델 매뉴얼과 대조', 'AJ65SBTB1 B RATE의 4=ON, 2=OFF, 1=OFF 설정을 해당 모델의 설정 기준과 대조'),
(201, '10 Mbps 허용 거리·국간 거리 및 배선 분리 기준을 제조사 매뉴얼로 확인', '10 Mbps 허용 거리·국간 거리 및 배선 분리를 제조사 설치 기준과 대조'),
(231, 'EMS SWITCH 01 / 02. 과거 캡처 X002D는 현재 의미와 다를 수 있음', 'EMS SWITCH 01 / 02'),
(254, '제공 사양서 V6.2: RS-485, 115200 bps, 8 bit, None parity, 1 stop bit', 'RS-485, 115200 bps, 8 bit, None parity, 1 stop bit'),
(291, '제공 캡처의 확인값 또는 주의', '설정 기준'),
(291, '192.168.100.238 설정 이력. 어느 리더인지 현장 Source/Target 명판과 대조', '각 리더에 승인된 고유 IP 설정'),
(291, '255.255.255.0 / 192.168.100.1 캡처. 실제 설비망 승인값 우선', '설비 네트워크 주소 계획에 따른 서브넷 마스크 및 Gateway'),
(291, 'Source와 다른 고유 IP 사용. 확인되지 않은 주소를 임의로 239 등으로 확정하지 않음', 'Source와 다른 고유 IP 사용'),
(323, 'W를 kW로 변환하는 책임은 PLC 규격과 합의. PC에는 kW ×1000 값 전달', 'PLC에서 kW 단위로 환산한 후 PC에 kW ×1000 값 전달'),
(323, 'AccuraSM은 해당 모델 설정용. PowerCableSimulator 값은 생산 실측값이 아님', '승인된 계측기 설정 소프트웨어 사용. 시뮬레이터 값을 생산 실측값으로 사용하지 않음'),
(351, '현재 코드 기준', '적용 기준'),
(374, '현재 코드의 색상/셀 조건 검사. 빈 슬롯 여부를 실제 확인', '채널 표시 상태와 셀 감지 조건 검사. 빈 슬롯 여부를 실제 확인'),
(381, '현재 동작과 주의', '동작과 주의'),
(432, '확인 가능한 캡처 이력', '관리 항목'),
(432, '최종 운영값 기록', '운영값 기록'),
(432, '운영 PC와 노트북이 다름', 'PC 식별 정보 및 운영체제'),
(432, 'Intel AMT SOL, COM8→COM3 이력', '드라이버 / COM 번호'),
(432, 'SR-X100W 설정 캡처 있음', 'Source 리더 식별 정보'),
(432, '별도 IP·장치 구분 필요', 'Target 리더 식별 정보'),
]
for index, old, new in cell_edits:
    cellreplace(index, old, new)

history = Table(els[16], doc._body)
for row in history.rows:
    for cell in row.cells:
        for par in cell.paragraphs:
            if '현행 코드 기준' in par.text:
                setpara(par, '설치·설정 및 유지보수 절차 신규 작성')
row = deepcopy(history.rows[-1]._tr)
for tc, value in zip(row.findall(qn('w:tc')), ['0.3', '2026.09.17', '본문 및 그림 설명 정비, 마지막 장 확인사항 통합', '대흥시스텍']):
    texts = tc.xpath('.//w:t')
    texts[0].text = value
    for tx in texts[1:]:
        tx.text = ''
history._tbl.append(row)

# Replace the old bibliography with a consolidated, customer-facing final chapter.
cursor = els[441].getprevious()
for el in els[441:]:
    if el.tag != qn('w:sectPr'):
        body.remove(el)

def append(el):
    global cursor
    cursor.addnext(el)
    cursor = el

title('부록 B 확인사항')
sub('B.1 설비 설정 및 인수 확인')
p('설비 인수, 부품 교체 및 설정 복원 시 아래 항목을 확인하고 결과를 부록 A에 기록합니다. 본문의 조작 절차에는 해당 설비의 승인된 운영값을 적용합니다.')
table(['항목', '확인사항'], [
['서보 백업과 운영값', '3장의 수치는 2026.08.13 백업 기준입니다. 이후 승인 변경이 있으면 최신 백업을 우선합니다. 이전 4축 화면의 5974h·1388h와 구분하고, 운영용 .prm2 경로 및 설치 프로그램의 기동 동작이 3.3절과 일치하는지 확인합니다.'],
['기계 단위와 센서', '전자기어 수치는 실제 이동량의 승인 기록을 대신하지 않습니다. 각 축의 1000 지령=1 mm, 방향·원점·DOG/리밋 극성을 실측·확인합니다. 자동계산의 45/590/95 mm는 계산 규칙이므로 실물 트레이 치수와 대조합니다.'],
['직렬 장치', '운영 PC에는 Intel AMT SOL이 설치되어 있으며 설치 화면에는 COM8에서 COM3로 변경한 이력이 있습니다. 장치명만으로 물리 RS-485 포트를 확정할 수 없으므로 하드웨어 ID, 실제 배선 포트 및 정상 수신을 함께 확인합니다.'],
['연기감지기', 'V6.2 통신 규격은 115200/8N1입니다. Human Automation 모드의 Baud 미지정 대체값은 9600이므로 프로그램에 115200을 명시합니다. 국번, PA1/PA2 경보값 및 PA4 프로토콜을 장치와 대조합니다.'],
['바코드 리더 2대', '설치 화면의 형명은 SR-X100W, 초기 검색 IP는 192.168.100.100입니다. 설정 화면에는 192.168.100.238, TCP 9004, 마스크 255.255.255.0, Gateway 192.168.100.1이 표시됩니다. 해당 주소의 Source/Target 대응과 두 리더의 최종 운영 IP는 명판·판독 시험으로 확인합니다.'],
['전력 계측', 'PC는 PLC 전력값을 사용합니다. Accura 2500/2550 자료는 PLC 측 계측 참고용이며 실제 설치 모델·CT/PT 비율·결선·PLC 취득 방식의 확정 기록은 아닙니다. UInt32 변환은 음수를 표현하지 않으므로 음수 전력 사용 여부도 확인합니다.'],
['보전 주기와 승인', '12.1절의 주기는 기본 점검 기준입니다. 청소·윤활·교체 주기는 제조사 지정, 생산 환경 및 사용량에 따라 보전 계획으로 확정합니다. 설치 모델의 지원 OS·드라이버와 교체품 적합성도 확인합니다.'],
], [2050, 7021])

p('B.2 화면 자료 및 적용 범위', 'h2', page=True)
p('화면과 사진은 조작 위치 및 점검 대상을 식별하기 위한 자료입니다. 화면의 숫자·출력 상태·시험 데이터는 아래 적용 범위에 따라 사용합니다.')
table(['구분', '확인사항'], [
['설치·프로그램 화면', 'Position Board, 시리얼카드 및 BCR 설치 화면은 운영 PC의 설정 이력입니다. Teaching·Configuration·Interface·Jog/Z 등 프로그램 화면에는 설명용 화면이 포함되어 있으며, 노트북 표시값과 운영 PC 설정은 다를 수 있습니다. 표시된 좌표·IP·온도·알람·시험 코드를 승인 운영값으로 복사하지 않습니다.'],
['CC-Link Master 화면', '그림 4-1은 국번호 1의 Standby master station 화면입니다. 운영 Master 설정 완료 상태를 나타내지 않으며, 실제 Master는 4.1절에 따라 국번호 0으로 확인합니다.'],
['CC-Link 오류 및 I/O 화면', '그림 4-3은 링크 오류 진단 화면입니다. 그림 4-4는 2026.08.06 상태 기록이며 Y 출력 조합을 정상 운전 기준으로 사용하지 않습니다. 이전 X002D 할당과 구분하여 현재 EMS SWITCH 02 및 최신 I/O 도면을 적용합니다.'],
['시리얼 설치 화면', '설치 화면의 드라이버 선택 이력은 다른 PC·시리얼 카드에 동일 드라이버를 적용하라는 의미가 아닙니다. 화면에 포함된 호환성 경고는 해당 하드웨어 ID와 제조사 지원 조건으로 판정합니다.'],
['Z 속도 표시 명칭', 'Z first 80% speed / Z final 20% speed는 고속·저속 설정의 화면 명칭입니다. 실제 전환 기준은 Target의 절대 Z 감속 시작 위치이며, 이동거리 비율 분할로 해석하지 않습니다.'],
['자료 및 검증 범위', '설치 자료, 운영 화면 및 프로그램 동작을 기준으로 정리했습니다. 서보 HOME 항목은 MR-MC200/300 IB-0300223 Control Parameters, 직렬 장치 식별은 Intel AMT SOL 자료와 함께 검토한 내용입니다. 현장 결선·안전회로·설치 파라미터의 최종 승인 기록 및 실제 동작·안전기능 시험을 대체하지 않습니다.'],
], [2050, 7021])
note('인수 및 생산 복귀 전 12.2절의 점검 결과, 적용 프로그램·설정 버전, 시험 일시와 승인자를 기록합니다.')

def heading(text):
    return next(par for par in doc.paragraphs if par.text == text and par.style.style_id == '10')

toc = heading('목 차')._p
el = toc.getnext()
while el is not None and el != heading('1. 안전 및 작업 준비')._p:
    nxt = el.getnext()
    body.remove(el)
    el = nxt
heads = [par.text for par in doc.paragraphs if par.style.style_id == '10' and par.text not in ('문서 정보', '목 차')]
pagefile = TMP/'toc_pages.json'
pagemap = json.loads(pagefile.read_text(encoding='utf-8')) if pagefile.exists() else {}
anchor = toc
for text in heads:
    el = deepcopy(ps[17])
    for c in list(el):
        if c.tag != qn('w:pPr'):
            el.remove(c)
    par = Paragraph(el, doc._body)
    par.paragraph_format.space_after = Pt(1)
    par.paragraph_format.space_before = Pt(0)
    par.add_run(text + '\t' + str(pagemap.get(text, '—')))
    anchor.addnext(el)
    anchor = el

doc.core_properties.revision = 3
buf = io.BytesIO()
doc.save(buf)
generated = ZipFile(io.BytesIO(buf.getvalue()))
editable = {'word/document.xml', 'docProps/core.xml'}
with ZipFile(OUT, 'w', ZIP_DEFLATED) as z:
    for name in original.namelist():
        z.writestr(name, generated.read(name) if name in editable else original.read(name))
with ZipFile(OUT) as z:
    assert set(z.namelist()) == set(original.namelist())
    for name in original.namelist():
        if name not in editable:
            assert z.read(name) == original.read(name), name
assert REF.read_bytes() == raw
(TMP/'headings.json').write_text(json.dumps(heads, ensure_ascii=False, indent=2), encoding='utf-8')
(TMP/'changes.json').write_text(json.dumps(changes, ensure_ascii=False, indent=2), encoding='utf-8')
(TMP/'artifact.md').write_text('''# Revision 003 contract
- Preserve reference 002; customer-submission editorial cleanup only.
- Preserve header/footer, logos, screenshots, relationships, styles, numbering and section geometry byte-for-byte.
- Move provenance, historic-value differences and unresolved hardware specifications into final appendix B.
- Retain operating behavior, parameter data, interlocks and safety warnings at point of use.
- Regenerate page-numbered TOC after Word pagination; render and inspect all pages.
- No program source, settings or hardware changes.
''', encoding='utf-8')
print(OUT)
print('PASS unchanged package parts; edits:', len(changes), 'headings:', len(heads))
